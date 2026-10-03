"""STEP 7 - Excel workbook with Summary, By Quarter and Run Log sheets."""

from __future__ import annotations

import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import EXCEL_OUT
from .logging_setup import get_logger

log = get_logger("output")

STATUS_FILL = {
    "OK": "C6EFCE",                    # green
    "PARTIAL": "FFEB9C",               # amber
    "NO TRANSCRIPTS FILED": "D9D9D9",  # grey
    "TRANSCRIPTS NOT ATTACHED": "BDD7EE",  # blue: exists, but on company site
    "SCRIP CODE NOT FOUND": "F8CBAD",  # orange
    "FAILED": "FFC7CE",                # red
}
STATUS_ORDER = {
    "OK": 0, "PARTIAL": 1, "TRANSCRIPTS NOT ATTACHED": 2,
    "NO TRANSCRIPTS FILED": 3, "SCRIP CODE NOT FOUND": 4, "FAILED": 5,
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _listify(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (list, tuple)):
        parts = []
        for item in v:
            if isinstance(item, dict):
                name = item.get("name", "")
                share = item.get("revenue_share", "")
                parts.append(f"{name} ({share})" if share else str(name))
            else:
                parts.append(str(item))
        return " | ".join(p for p in parts if p)
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False)
    return str(v)


def _style_header(ws, widths: dict[str, int], wrap_cols: set[str] | None = None) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    if wrap_cols:
        for col in wrap_cols:
            for cell in ws[col][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")


def _base_status(status: str) -> str:
    return "PARTIAL" if status.startswith("PARTIAL") else status


def write_excel(companies: list[dict], per_quarter: list[dict],
                runlog: list[dict], path=EXCEL_OUT) -> None:
    wb = Workbook()

    # ------------------------------------------------------------- Summary
    ws = wb.active
    ws.title = "Summary"
    cols = ["symbol", "scrip_code", "status", "transcripts_found", "quarters_covered",
            "shift_score", "narrative_shift", "new_in_latest", "dropped",
            "tone_trajectory"]
    ws.append(cols)

    def sort_key(row: dict):
        base = _base_status(row.get("status", "FAILED"))
        # shift_score is blank when the LLM stage was skipped or had too little
        # data, so coerce rather than assuming an int.
        try:
            score = int(row.get("shift_score"))
        except (TypeError, ValueError):
            score = -1
        return (STATUS_ORDER.get(base, 9), -score, row.get("symbol", ""))

    for row in sorted(companies, key=sort_key):
        ws.append([
            row.get("symbol", ""),
            row.get("scrip_code", ""),
            row.get("status", ""),
            row.get("transcripts_found", 0),
            row.get("quarters_covered", ""),
            row.get("shift_score", ""),
            _listify(row.get("narrative_shift")),
            _listify(row.get("new_in_latest")),
            _listify(row.get("dropped")),
            _listify(row.get("tone_trajectory")),
        ])

    for r in range(2, ws.max_row + 1):
        cell = ws.cell(row=r, column=3)
        colour = STATUS_FILL.get(_base_status(str(cell.value or "")))
        if colour:
            cell.fill = PatternFill("solid", fgColor=colour)

    _style_header(
        ws,
        {"A": 14, "B": 12, "C": 22, "D": 17, "E": 26, "F": 12,
         "G": 60, "H": 45, "I": 40, "J": 32},
        wrap_cols={"G", "H", "I", "J", "E"},
    )

    # ---------------------------------------------------------- By Quarter
    ws2 = wb.create_sheet("By Quarter")
    cols2 = ["symbol", "quarter", "news_dt", "kind", "n_pages", "n_chars",
             "extract_method", "business_summary", "key_numbers", "what_changed",
             "new_themes", "segments", "order_book", "capex", "guidance",
             "margin_commentary", "demand_commentary", "management_tone",
             "tone_justification", "notable_quotes", "risks_flagged"]
    ws2.append(cols2)
    for row in per_quarter:
        ws2.append([_listify(row.get(c)) if c in
                    ("segments", "new_themes", "notable_quotes", "risks_flagged",
                     "key_numbers", "what_changed")
                    else row.get(c, "") for c in cols2])
    _style_header(
        ws2,
        {"A": 12, "B": 10, "C": 20, "D": 20, "E": 9, "F": 10, "G": 14,
         "H": 45, "I": 45, "J": 45, "K": 38, "L": 32, "M": 26, "N": 30,
         "O": 34, "P": 40, "Q": 36, "R": 14, "S": 34, "T": 60, "U": 38},
        wrap_cols={"H","I","J","K","L","M","N","O","P","Q","S","T","U"},
    )

    # ------------------------------------------------------------- Run Log
    ws3 = wb.create_sheet("Run Log")
    ws3.append(["symbol", "quarter", "stage", "detail", "timestamp"])
    for e in runlog:
        ws3.append([e.get("symbol", ""), e.get("quarter", ""), e.get("stage", ""),
                    e.get("detail", ""), e.get("ts", "")])
    _style_header(ws3, {"A": 12, "B": 10, "C": 22, "D": 90, "E": 20},
                  wrap_cols={"D"})

    try:
        wb.save(path)
        log.info("wrote %s", path)
    except PermissionError:
        alt = path.with_name(path.stem + "_new.xlsx")
        wb.save(alt)
        log.warning("%s was locked, wrote %s instead", path.name, alt.name)
